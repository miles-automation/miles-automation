import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.automations.pdf_batch import (
    AutomationError,
    PageText,
    extract_field,
    run_pdf_batch,
)
from backend.automations.spec import load_pdf_spec

SPECS = Path(__file__).parents[1] / "gig_specs"


def _property_spec():
    return load_pdf_spec(SPECS / "pdf/property-records.json")


def test_extract_field_preserves_page_provenance_and_normalizes() -> None:
    spec = load_pdf_spec(SPECS / "pdf/fuel-delivery-invoices.json")
    date_field = next(field for field in spec.fields if field.name == "delivery_date")
    result = extract_field(
        [PageText(1, "Invoice No: A-10"), PageText(2, "Delivery Date: 07/20/2026")],
        date_field,
    )
    assert result.value == "2026-07-20"
    assert result.status == "ok"
    assert result.source_pages == [2]


def test_distinct_values_are_flagged_as_ambiguous() -> None:
    field = _property_spec().fields[0]
    result = extract_field([PageText(1, "Parcel ID: ABC-1\nParcel ID: ABC-2")], field)
    assert result.status == "ambiguous"
    assert result.value is None
    assert "ABC-1" in (result.detail or "")


def test_run_writes_delivery_workbook_and_audit_manifest(tmp_path: Path) -> None:
    incoming = tmp_path / "incoming"
    outgoing = tmp_path / "outgoing"
    incoming.mkdir()
    source = incoming / "record-001.pdf"
    source.write_bytes(b"fixture bytes")

    def fake_reader(path: Path) -> list[PageText]:
        assert path == source
        return [
            PageText(
                1,
                "\n".join(
                    [
                        "Parcel ID: 12-345-678",
                        "Property Address: 10 Main St",
                        "City: Royal Oak",
                        "Interested Parties: Jane Doe; Acme Bank",
                    ]
                ),
            )
        ]

    result = run_pdf_batch(_property_spec(), incoming, outgoing, reader=fake_reader)
    assert result.documents == 1
    assert result.pages == 1
    assert result.acceptance_passed is True
    assert result.required_exception_rate == 0.0

    workbook = load_workbook(result.workbook_path)
    sheet = workbook["Property Records"]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, [cell.value for cell in sheet[2]], strict=True))
    assert row["parcel_id"] == "12-345-678"
    assert row["parcel_id__source_pages"] == "1"
    assert row["status"] == "ok"
    assert workbook["Exceptions"].max_row == 1

    manifest = json.loads(result.manifest_path.read_text())
    assert manifest["summary"]["acceptance_passed"] is True
    assert manifest["documents"][0]["source_sha256"]
    assert manifest["spec"]["fingerprint"] == _property_spec().fingerprint


def test_required_exception_fails_acceptance_but_still_delivers(tmp_path: Path) -> None:
    incoming = tmp_path / "incoming"
    incoming.mkdir()
    (incoming / "record.pdf").write_bytes(b"fixture")

    result = run_pdf_batch(
        _property_spec(),
        incoming,
        tmp_path / "outgoing",
        reader=lambda _path: [
            PageText(1, "Parcel ID: 12-345\nDocument notes: enough text")
        ],
    )
    assert result.acceptance_passed is False
    assert result.required_exception_rate == 0.75
    workbook = load_workbook(result.workbook_path)
    assert workbook["Exceptions"].max_row == 4


def test_document_limit_is_enforced_before_extraction(tmp_path: Path) -> None:
    incoming = tmp_path / "incoming"
    incoming.mkdir()
    (incoming / "one.pdf").write_bytes(b"one")
    (incoming / "two.pdf").write_bytes(b"two")
    payload = _property_spec().model_copy(deep=True)
    payload.limits.max_documents = 1

    with pytest.raises(AutomationError, match="batch has 2 documents"):
        run_pdf_batch(
            payload,
            incoming,
            tmp_path / "outgoing",
            reader=lambda _path: [PageText(1, "")],
        )


def test_image_only_pdf_is_rejected_for_ocr_or_review(tmp_path: Path) -> None:
    incoming = tmp_path / "incoming"
    incoming.mkdir()
    (incoming / "scan.pdf").write_bytes(b"scan")

    with pytest.raises(AutomationError, match="image-only scans need OCR"):
        run_pdf_batch(
            _property_spec(),
            incoming,
            tmp_path / "outgoing",
            reader=lambda _path: [PageText(1, "")],
        )


def test_workbook_escapes_formula_like_document_values(tmp_path: Path) -> None:
    incoming = tmp_path / "incoming"
    incoming.mkdir()
    (incoming / "record.pdf").write_bytes(b"fixture")

    result = run_pdf_batch(
        _property_spec(),
        incoming,
        tmp_path / "outgoing",
        reader=lambda _path: [
            PageText(
                1,
                "\n".join(
                    [
                        "Parcel ID: 12-345",
                        'Address: =HYPERLINK("https://example.test")',
                        "City: Royal Oak",
                        "Interested Parties: Jane Doe",
                    ]
                ),
            )
        ],
    )
    workbook = load_workbook(result.workbook_path, data_only=False)
    sheet = workbook["Property Records"]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, [cell.value for cell in sheet[2]], strict=True))
    assert row["address"].startswith("'=HYPERLINK")
