from __future__ import annotations

import hashlib
import json
import re
import time
import uuid
from collections.abc import Callable, Sequence
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from pypdf import PdfReader

from backend.automations.spec import ExtractionField, PdfGigSpec

FieldStatus = Literal["ok", "missing", "ambiguous", "invalid"]
InputStatus = Literal["text_ready", "ocr_or_manual_review"]


class AutomationError(RuntimeError):
    pass


@dataclass(frozen=True)
class PageText:
    page_number: int
    text: str


@dataclass(frozen=True)
class FieldResult:
    name: str
    value: str | None
    status: FieldStatus
    source_pages: list[int]
    detail: str | None = None


@dataclass(frozen=True)
class DocumentResult:
    source_file: str
    source_sha256: str
    page_count: int
    input_status: InputStatus
    status: Literal["ok", "exceptions"]
    elapsed_ms: int
    fields: list[FieldResult]


@dataclass(frozen=True)
class RunResult:
    run_id: str
    workbook_path: Path
    manifest_path: Path
    documents: int
    pages: int
    layouts: int
    required_exception_rate: float
    acceptance_passed: bool
    elapsed_ms: int


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_pdf_pages(path: Path, max_pages: int) -> list[PageText]:
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise AutomationError(f"encrypted PDF is out of scope: {path.name}")
    if len(reader.pages) > max_pages:
        raise AutomationError(
            f"{path.name} has {len(reader.pages)} pages; limit is {max_pages}"
        )
    return [
        PageText(page_number=index, text=page.extract_text() or "")
        for index, page in enumerate(reader.pages, start=1)
    ]


def _normalize(raw: str, field: ExtractionField) -> tuple[str | None, str | None]:
    value = " ".join(raw.strip().split())
    try:
        if field.value_type == "integer":
            value = str(int(value.replace(",", "")))
        elif field.value_type == "decimal":
            decimal = Decimal(re.sub(r"[$,\s]", "", value))
            value = format(decimal, "f")
        elif field.value_type == "date":
            parsed = None
            for date_format in field.date_formats:
                try:
                    parsed = datetime.strptime(value, date_format).date()
                    break
                except ValueError:
                    continue
            if parsed is None:
                return None, f"unsupported date: {value}"
            value = parsed.isoformat()
        elif field.value_type == "email":
            value = value.lower()
            if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
                return None, f"invalid email: {value}"
    except (ValueError, InvalidOperation) as exc:
        return None, f"invalid {field.value_type}: {value} ({exc})"

    if field.validation_pattern and not re.fullmatch(field.validation_pattern, value):
        return None, f"value failed validation: {value}"
    return value, None


def extract_field(pages: list[PageText], field: ExtractionField) -> FieldResult:
    candidates: list[tuple[str, int]] = []
    errors: list[str] = []
    for page in pages:
        for pattern in field.patterns:
            for match in re.finditer(pattern, page.text, re.IGNORECASE | re.MULTILINE):
                normalized, error = _normalize(match.group("value"), field)
                if error:
                    errors.append(f"page {page.page_number}: {error}")
                elif normalized is not None:
                    candidates.append((normalized, page.page_number))

    unique_values = {value for value, _ in candidates}
    source_pages = sorted({page for _, page in candidates})
    if len(unique_values) > 1:
        values = ", ".join(sorted(unique_values))
        return FieldResult(
            field.name,
            None,
            "ambiguous",
            source_pages,
            f"multiple distinct values: {values}",
        )
    if len(unique_values) == 1 and errors:
        return FieldResult(
            field.name,
            None,
            "ambiguous",
            source_pages,
            "; ".join(errors),
        )
    if len(unique_values) == 1:
        return FieldResult(field.name, next(iter(unique_values)), "ok", source_pages)
    if errors:
        return FieldResult(field.name, None, "invalid", [], "; ".join(errors))
    return FieldResult(
        field.name,
        None,
        "missing",
        [],
        "required value not found" if field.required else "optional value not found",
    )


def extract_document(
    path: Path,
    spec: PdfGigSpec,
    *,
    reader: Callable[[Path, int], list[PageText]] = read_pdf_pages,
) -> DocumentResult:
    started = time.perf_counter()
    pages = reader(path, spec.limits.max_pages_per_document)
    if len(pages) > spec.limits.max_pages_per_document:
        raise AutomationError(
            f"{path.name} has {len(pages)} pages; limit is "
            f"{spec.limits.max_pages_per_document}"
        )
    extracted_characters = sum(len(page.text.strip()) for page in pages)
    input_status: InputStatus = (
        "text_ready"
        if extracted_characters >= spec.limits.min_extracted_text_characters
        else "ocr_or_manual_review"
    )
    fields = [extract_field(pages, field) for field in spec.fields]
    required = {field.name for field in spec.fields if field.required}
    has_exceptions = (
        input_status != "text_ready"
        or any(result.status != "ok" and result.name in required for result in fields)
        or any(result.status in {"ambiguous", "invalid"} for result in fields)
    )
    return DocumentResult(
        source_file=path.name,
        source_sha256=_sha256(path),
        page_count=len(pages),
        input_status=input_status,
        status="exceptions" if has_exceptions else "ok",
        elapsed_ms=round((time.perf_counter() - started) * 1000),
        fields=fields,
    )


def _style_sheet(sheet: Worksheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        sheet.column_dimensions[column[0].column_letter].width = width


def _excel_safe(
    value: str | int | float | bool | None,
) -> str | int | float | bool | None:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        if value.startswith(("+", "-")):
            try:
                Decimal(value)
            except InvalidOperation:
                pass
            else:
                return value
        return f"'{value}"
    return value


def _append_safe(
    sheet: Worksheet, values: Sequence[str | int | float | bool | None]
) -> None:
    sheet.append([_excel_safe(value) for value in values])


def _write_workbook(
    path: Path,
    spec: PdfGigSpec,
    documents: list[DocumentResult],
    *,
    run_id: str,
    layout_count: int,
    required_exception_rate: float,
    acceptance_passed: bool,
) -> None:
    workbook = Workbook()
    data = workbook.active
    assert isinstance(data, Worksheet)
    data.title = spec.output.workbook_sheet
    headers = [
        "source_file",
        "source_sha256",
        "page_count",
        "input_status",
        "status",
    ]
    for extraction_field in spec.fields:
        headers.append(extraction_field.name)
        if spec.output.include_source_page:
            headers.append(f"{extraction_field.name}__source_pages")
        if spec.output.include_field_status:
            headers.append(f"{extraction_field.name}__status")
    _append_safe(data, headers)

    for document in documents:
        fields = {result.name: result for result in document.fields}
        row: list[str | int | None] = [
            document.source_file,
            document.source_sha256,
            document.page_count,
            document.input_status,
            document.status,
        ]
        for field_spec in spec.fields:
            result = fields[field_spec.name]
            row.append(result.value)
            if spec.output.include_source_page:
                row.append(",".join(str(page) for page in result.source_pages))
            if spec.output.include_field_status:
                row.append(result.status)
        _append_safe(data, row)
    _style_sheet(data)

    exceptions = workbook.create_sheet("Exceptions")
    _append_safe(
        exceptions,
        ["source_file", "field", "status", "value", "source_pages", "detail"],
    )
    for document in documents:
        for result in document.fields:
            if result.status != "ok":
                _append_safe(
                    exceptions,
                    [
                        document.source_file,
                        result.name,
                        result.status,
                        result.value,
                        ",".join(str(page) for page in result.source_pages),
                        result.detail,
                    ],
                )
    _style_sheet(exceptions)

    summary = workbook.create_sheet("Run Summary")
    _append_safe(summary, ["metric", "value"])
    _append_safe(summary, ["run_id", run_id])
    _append_safe(summary, ["spec_slug", spec.slug])
    _append_safe(summary, ["spec_fingerprint", spec.fingerprint])
    _append_safe(summary, ["documents", len(documents)])
    _append_safe(summary, ["pages", sum(document.page_count for document in documents)])
    _append_safe(summary, ["layouts", layout_count])
    _append_safe(summary, ["required_exception_rate", required_exception_rate])
    _append_safe(summary, ["acceptance_passed", acceptance_passed])
    _style_sheet(summary)
    workbook.save(path)


def run_pdf_batch(
    spec: PdfGigSpec,
    input_dir: str | Path,
    output_dir: str | Path,
    *,
    layout_count: int,
    reader: Callable[[Path, int], list[PageText]] = read_pdf_pages,
) -> RunResult:
    started = time.perf_counter()
    started_at = datetime.now(timezone.utc)
    source_dir = Path(input_dir)
    if layout_count < 1 or layout_count > spec.limits.max_layouts:
        raise AutomationError(
            f"batch declares {layout_count} layouts; limit is {spec.limits.max_layouts}"
        )
    files = sorted(
        path for path in source_dir.iterdir() if path.suffix.lower() == ".pdf"
    )
    if not files:
        raise AutomationError(f"no PDF files found in {source_dir}")
    if len(files) > spec.limits.max_documents:
        raise AutomationError(
            f"batch has {len(files)} documents; limit is {spec.limits.max_documents}"
        )

    documents: list[DocumentResult] = []
    total_pages = 0
    for path in files:
        document = extract_document(path, spec, reader=reader)
        total_pages += document.page_count
        if total_pages > spec.limits.max_total_pages:
            raise AutomationError(
                f"batch has more than {spec.limits.max_total_pages} total pages"
            )
        documents.append(document)

    required_names = {field.name for field in spec.fields if field.required}
    required_cells = len(documents) * len(required_names)
    required_exceptions = sum(
        result.status != "ok"
        for document in documents
        for result in document.fields
        if result.name in required_names
    )
    required_exception_rate = (
        round(required_exceptions / required_cells, 6) if required_cells else 0.0
    )
    acceptance_passed = (
        all(document.input_status == "text_ready" for document in documents)
        and required_exception_rate <= spec.acceptance.max_required_exception_rate
    )

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    run_id = f"{started_at:%Y%m%dT%H%M%SZ}-{uuid.uuid4().hex[:8]}"
    workbook_path = destination / f"{spec.slug}-{run_id}.xlsx"
    manifest_path = destination / f"{spec.slug}-{run_id}.audit.json"
    _write_workbook(
        workbook_path,
        spec,
        documents,
        run_id=run_id,
        layout_count=layout_count,
        required_exception_rate=required_exception_rate,
        acceptance_passed=acceptance_passed,
    )

    elapsed_ms = round((time.perf_counter() - started) * 1000)
    manifest = {
        "schema_version": 1,
        "run_id": run_id,
        "started_at": started_at.isoformat(),
        "completed_at": datetime.now(timezone.utc).isoformat(),
        "elapsed_ms": elapsed_ms,
        "spec": {
            "slug": spec.slug,
            "fingerprint": spec.fingerprint,
            "source": spec.source.model_dump(mode="json"),
        },
        "summary": {
            "documents": len(documents),
            "pages": total_pages,
            "layouts": layout_count,
            "required_exceptions": required_exceptions,
            "required_exception_rate": required_exception_rate,
            "acceptance_passed": acceptance_passed,
        },
        "documents": [asdict(document) for document in documents],
        "outputs": {"workbook": workbook_path.name},
    }
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")

    return RunResult(
        run_id=run_id,
        workbook_path=workbook_path,
        manifest_path=manifest_path,
        documents=len(documents),
        pages=total_pages,
        layouts=layout_count,
        required_exception_rate=required_exception_rate,
        acceptance_passed=acceptance_passed,
        elapsed_ms=elapsed_ms,
    )
